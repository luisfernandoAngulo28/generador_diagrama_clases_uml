"""
Construye el documento final del 2do Parcial (Word) a partir de las secciones
Markdown de docs/. Uso (desde la raíz del repo):

    python docs/documento_final/build_documento.py

Requiere: pandoc y python-docx. Los QR se generan con generar_qr.py.
Al terminar, abrir el .docx en Word y pulsar F9 sobre el índice si no se
actualizó solo (o usar finalizar_con_word.ps1).
"""
from __future__ import annotations

import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / "docs"
OUT_DIR = DOCS / "documento_final"
BUILD = OUT_DIR / "_build"
UML_DIR = DOCS / "diagrams" / "2do_parcial"
OUTPUT = OUT_DIR / "Segundo_Parcial_Grupo43_SW1.docx"

SECCIONES = [
    DOCS / "seccion1_fundamentacion_teorica_2do_parcial.md",
    DOCS / "seccion2_registro_del_desarrollo.md",
    DOCS / "seccion3_aprendizaje_del_usuario.md",
]
ANEXOS = DOCS / "seccion_anexos.md"

PAGE_BREAK = '```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```\n'
TOC_FIELD = (
    "```{=openxml}\n"
    '<w:p><w:pPr><w:pStyle w:val="TOCHeading"/></w:pPr><w:r><w:t>Índice</w:t></w:r></w:p>'
    '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
    '<w:r><w:instrText xml:space="preserve"> TOC \\o "1-2" \\h \\z \\u </w:instrText></w:r>'
    '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
    "<w:r><w:t>(Actualizar el índice: clic derecho → Actualizar campo)</w:t></w:r>"
    '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>\n'
    "```\n"
)

DIAGRAMAS = [
    ("01_despliegue.png", "Diagrama de despliegue (nodos, entornos de ejecución, artefactos y rutas de comunicación)",
     "docker-compose.prod.yml, docker-compose.local.yml, frontend/nginx.conf"),
    ("02_componentes.png", "Diagrama de componentes (interfaces REST y Socket.io provistas y requeridas)",
     "backend/src/*/…module.ts, frontend/src/api, flutter-template/lib/live"),
    ("03_clases_persistencia.png", "Diagrama de clases — persistencia políglota con autoría granular embebida",
     "backend/src/diagrams/schemas, uml.types.ts, entidades de PostgreSQL"),
    ("04_clases_motor_ia.png", "Diagrama de clases — motor de IA híbrido (patrón Strategy) y Moderador grupal",
     "backend/src/ai/**"),
    ("05_secuencia_chat_ia.png", "Diagrama de secuencia — pregunta pública al Moderador IA (alt / opt / par)",
     "ChatPanel.tsx, chat.gateway.ts, moderator.service.ts, hybrid.provider.ts"),
    ("06_secuencia_guardado_autoria.png", "Diagrama de secuencia — guardado con autoría granular y bitácora",
     "App.tsx (saveDiagram), diagrams.controller.ts, diagrams.service.ts, authorship.util.ts"),
    ("07_estados_sesion.png", "Máquina de estados — sesión continua (ACTIVE ⇄ PAUSED)",
     "sessions.service.ts, diagrams.gateway.ts"),
    ("08_actividad_motor_hibrido.png", "Diagrama de actividad — decisión del motor de IA híbrido",
     "hybrid.provider.ts, ai.service.ts"),
]


def caratula() -> str:
    qr_repo = (OUT_DIR / "img" / "qr_repositorio.png").as_posix()
    qr_prod = (OUT_DIR / "img" / "qr_produccion.png").as_posix()
    return f"""
::: {{custom-style="Cover Institution"}}
⟨Universidad · Facultad · Carrera⟩
:::

::: {{custom-style="Cover Subject"}}
Ingeniería de Software 1 (SW1)
:::

::: {{custom-style="Cover Title"}}
SEGUNDO EXAMEN PARCIAL
:::

::: {{custom-style="Cover App"}}
Diagramador UML
:::

::: {{custom-style="Cover Subtitle"}}
Herramienta CASE colaborativa con Moderador IA para diagramas de clases UML 2.5 y generación de backends Spring Boot
:::

::: {{custom-style="Cover Info"}}
**Grupo:** 43

**Integrantes:**

Elizabeth Rivero Peña

Luis Fernando Angulo Heredia

**Docente:** Ing. Rolando Martínez Canedo

**Fecha de entrega:** miércoles 11 de noviembre de 2026
:::

| ![]({qr_repo}){{width=1.45in}} | ![]({qr_prod}){{width=1.45in}} |
|:---:|:---:|
| **Repositorio (GitHub)** | **Aplicación en producción** |

{PAGE_BREAK}
{TOC_FIELD}
"""


def limpiar_seccion(md: str) -> str:
    """Quita la nota editorial inicial (blockquote tras el H1) y marca los
    pedidos de captura como recuadros a completar."""
    lines = md.splitlines()
    out: list[str] = []
    i = 0
    # H1 y posible blockquote inmediatamente posterior
    while i < len(lines):
        out.append(lines[i])
        if lines[i].startswith("# "):
            i += 1
            while i < len(lines) and (lines[i].strip() == "" or lines[i].startswith(">")):
                i += 1
            out.append("")
            break
        i += 1
    out.extend(lines[i:])
    text = "\n".join(out)
    # Separadores horizontales de Markdown: en Word sobran
    text = re.sub(r"^---\s*$", "", text, flags=re.M)
    # *(Captura N: …)* / *(Figura A.n: …)* (puede ocupar varias líneas) → recuadro amarillo
    text = re.sub(
        r"^\*\((Captura|Figura)(.*?)\)\*[ \t]*$",
        lambda m: '::: {custom-style="Placeholder"}\n[Insertar '
        + m.group(1).lower()
        + " ".join(m.group(2).split())
        + "]\n:::",
        text,
        flags=re.M | re.S,
    )
    return separar_listas(text)


LIST_ITEM = re.compile(r"^(\s*)([-*+]|\d+\.)\s+\S")


def separar_listas(text: str) -> str:
    """Pandoc solo reconoce una lista si va precedida de una línea en blanco;
    en los .md hay listas pegadas a un párrafo ("**Evidencia:**" + "- …")."""
    out: list[str] = []
    in_code = False
    for line in text.splitlines():
        if line.lstrip().startswith("```"):
            in_code = not in_code
        prev = out[-1] if out else ""
        item = LIST_ITEM.match(line)
        if (
            not in_code
            and item
            and not item.group(1)  # solo ítems de primer nivel
            and prev.strip()
            and not LIST_ITEM.match(prev)
            and not prev.startswith((" ", "\t", "|", ">", "#"))
        ):
            out.append("")
        out.append(line)
    return "\n".join(out)


def seccion_uml() -> str:
    parts = [
        "## Anexo E. Modelos UML 2.5.1 del 2do Parcial",
        "",
        "Los siguientes modelos usan notación **UML 2.5.1** (Object Management Group, 2017) "
        "y se construyeron a partir del código real. Debajo de cada figura se indica qué "
        "archivos del código refleja: es la evidencia de V&V de que *lo que está en el "
        "código está reflejado en los modelos*. Las fuentes PlantUML están en "
        "`docs/diagrams/2do_parcial/`.",
        "",
    ]
    for n, (png, titulo, codigo) in enumerate(DIAGRAMAS, start=1):
        path, width_in = preparar_figura(UML_DIR / png)
        parts += [
            PAGE_BREAK,
            f"![Figura E.{n}. {titulo}]({path.as_posix()}){{width={width_in:.2f}in}}",
            "",
            f"*Refleja el código en:* `{codigo}`",
            "",
        ]
    return "\n".join(parts)


# Caja disponible para una figura en página carta con márgenes de 2,5 cm,
# dejando lugar al pie de figura y a la línea de trazabilidad.
BOX_W_IN, BOX_H_IN = 6.4, 7.3


def preparar_figura(png: Path) -> tuple[Path, float]:
    """Los diagramas apaisados se rotan 90° para aprovechar el alto de la
    hoja (se leen girando la página). Devuelve la ruta y el ancho en pulgadas."""
    from PIL import Image

    img = Image.open(png)
    w, h = img.size
    if w / h > 1.3:
        img = img.rotate(90, expand=True)
        out = BUILD / f"rot_{png.name}"
        img.save(out)
        w, h = img.size
        png = out
    width_in = min(BOX_W_IN, BOX_H_IN * w / h)
    return png, width_in


def construir_markdown() -> Path:
    BUILD.mkdir(parents=True, exist_ok=True)
    chunks = [caratula()]
    for path in SECCIONES:
        chunks.append(limpiar_seccion(path.read_text(encoding="utf-8")))
    anexos = limpiar_seccion(ANEXOS.read_text(encoding="utf-8"))
    # Insertar el Anexo E (UML) antes de la bibliografía (Anexo F, al final)
    marker = "## Anexo F."
    anexos = anexos.replace(marker, seccion_uml() + "\n\n" + marker, 1)
    chunks.append(anexos)
    master = BUILD / "documento.md"
    master.write_text("\n\n".join(chunks), encoding="utf-8")
    return master


def _set_font(style, size: float | None = None, bold: bool | None = None, color: RGBColor | None = None):
    style.font.name = "Arial"
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), "Arial")
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color


def construir_referencia() -> Path:
    """Plantilla de estilos para pandoc: Arial, títulos azul oscuro, carátula y recuadros."""
    ref = BUILD / "referencia.docx"
    subprocess.run(
        ["pandoc", "-o", str(ref), "--print-default-data-file", "reference.docx"],
        check=True,
    )
    doc = Document(ref)
    azul = RGBColor(0x1F, 0x3A, 0x5F)
    styles = doc.styles

    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in [s.name for s in styles]:
            _set_font(styles[name], 11)
            pf = styles[name].paragraph_format
            pf.space_after = Pt(6)
            pf.line_spacing = 1.15
    for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12), ("Heading 4", 11)):
        _set_font(styles[name], size, True, azul)
        styles[name].paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        styles[name].paragraph_format.keep_with_next = True
    # Cada sección (Heading 1) empieza en página nueva.
    styles["Heading 1"].paragraph_format.page_break_before = True

    def para_style(name: str, size: float, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.CENTER,
                   before=0, after=6):
        st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH) if name not in [s.name for s in styles] else styles[name]
        st.base_style = styles["Normal"]
        _set_font(st, size, bold, color)
        st.paragraph_format.alignment = align
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        return st

    para_style("Cover Institution", 13, True, azul, before=24, after=4)
    para_style("Cover Subject", 13, False, azul, after=60)
    para_style("Cover Title", 26, True, azul, after=18)
    para_style("Cover App", 20, True, RGBColor(0x2C, 0x3E, 0x50), after=6)
    para_style("Cover Subtitle", 12, False, RGBColor(0x55, 0x55, 0x55), after=48)
    para_style("Cover Info", 12, False, None, after=2)
    para_style("Cover QR", 11, False, None, before=36, after=0)
    para_style("Cover QR Caption", 10, True, azul, after=0)
    ph = para_style("Placeholder", 10, False, RGBColor(0x7A, 0x5C, 0x00), align=WD_ALIGN_PARAGRAPH.LEFT,
                    before=4, after=8)
    ppr = ph.element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "FFF3C4")
    ppr.append(shd)
    if "TOC Heading" in [s.name for s in styles]:
        _set_font(styles["TOC Heading"], 18, True, azul)
    if "Image Caption" in [s.name for s in styles]:
        _set_font(styles["Image Caption"], 10, False, RGBColor(0x44, 0x44, 0x44))
        styles["Image Caption"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if "Source Code" in [s.name for s in styles]:
        styles["Source Code"].font.size = Pt(7)

    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.59), Cm(27.94)  # Carta
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Cm(2.5))
    section.top_margin = section.bottom_margin = Cm(2.5)
    doc.save(ref)
    return ref


def _borde(tbl):
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "9AA5B1")
        borders.append(el)
    tblPr.append(borders)


def post_procesar(path: Path):
    """Bordes y encabezado sombreado en las tablas, y tamaño de letra compacto."""
    doc = Document(path)
    # Bloques de código (diagramas ASCII): letra chica para que no se corten.
    for p in doc.paragraphs:
        if p.style is not None and p.style.name == "Source Code":
            for run in p.runs:
                run.font.size = Pt(7)
    for tbl in doc.tables:
        if tbl._tbl.xpath(".//w:drawing"):
            # Tabla de la carátula (QR): sin bordes ni sombreado, centrada.
            for row in tbl.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for run in p.runs:
                            run.font.size = Pt(10)
                            run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
            continue
        _borde(tbl)
        for i, row in enumerate(tbl.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9.5)
                        if i == 0:
                            run.font.bold = True
                if i == 0:
                    tcPr = cell._tc.get_or_add_tcPr()
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), "D6EAF8")
                    tcPr.append(shd)
    doc.save(path)


def main():
    master = construir_markdown()
    ref = construir_referencia()
    subprocess.run(
        [
            "pandoc", str(master),
            "-f", "markdown+fenced_divs+raw_attribute+pipe_tables+task_lists",
            "-t", "docx",
            "--reference-doc", str(ref),
            "--resource-path", str(ROOT),
            "-o", str(OUTPUT),
        ],
        check=True,
    )
    post_procesar(OUTPUT)
    print(f"Documento generado: {OUTPUT}")


if __name__ == "__main__":
    main()
